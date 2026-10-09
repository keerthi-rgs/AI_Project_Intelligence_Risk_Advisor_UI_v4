"""Robustness test: agents must extract scope from documents that do NOT follow the sample-data layout.

Builds a DOCX with Word headings + a roles table, a TXT course-style milestone plan, a CSV task tracker,
and a PDF status report, then checks goals / milestones / deliverables / responsibilities, risks, and blockers.
"""
import io
import os
import sys

sys.path.insert(0, os.path.abspath("."))

from docx import Document

from app import extract_docx, extract_txt, extract_csv, extract_pdf
from agents import ScopeDeliverableAgent, RiskForecastingAgent, BlockerActionItemAgent

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

PLAN_TXT = """AI Project Intelligence & Risk Advisor

Project Objectives
- Help project teams understand scope, risks and blockers directly from their documents.
- Provide a grounded conversational assistant over uploaded project content.

Milestone 3 (Week 5-6 | ~10 Hours)
1. Develop the risk forecasting agent that flags schedule slippage and dependency gaps.
2. Build blocker and action item extraction from meeting notes.

Milestone 4 (Week 7-8 | ~10 Hours)
1. Build Project Insights and Risk Summary Dashboard — displays health score, extracted risks, action items, and scope summary in a clean visual layout.
2. Implement incremental document upload — teams can add new meeting notes or progress updates and the knowledge base updates automatically.
3. Conduct end-to-end testing across varied document types, agent outputs, and conversational workflows.
4. Prepare technical documentation, project report, and final demonstration.

Evaluation Criteria
1. Accuracy and completeness of scope extraction, risk detection, and blocker identification from uploaded documents.
2. Quality and usefulness of auto-generated user stories, risk registers, and action item lists.
"""

CSV_TRACKER = """Task,Type,Owner,Deadline,Status
Requirements sign-off,Milestone,Priya Nair,2026-11-05,Done
Beta release,Milestone,Ravi Kumar,2026-12-01,In Progress
Requirements specification document,Deliverable,Priya Nair,2026-11-01,Done
Test plan and test report,Deliverable,Anita Rao,2026-11-25,Open
Fix login timeout bug,Task,Ravi Kumar,2026-11-10,Blocked
"""

MEETING_TXT = """Team Sync - 3 Nov 2026
Discussion
- The login service keeps timing out under load, so the team is unable to run performance tests.
Decisions needed
- Whether to move the public launch to January.
Next steps
- Ravi will fix the login timeout by 10 Nov 2026.
- Priya to send the updated test plan to QA.
- Update the project risk log before Friday's review.
"""

PDF_LINES = [
    "Weekly Status Report - Week 6",
    "The aim of this release is to deliver the reporting module to pilot users.",
    "Milestone: Pilot launch by 20 Nov 2026.",
    "The QA lead is responsible for regression testing of the reporting module.",
    "We are blocked by the vendor API outage; integration testing cannot start.",
    "There is a risk of schedule slippage because the data team is understaffed.",
]


def make_docx() -> bytes:
    d = Document()
    d.add_heading("Customer Portal Redesign - Project Plan", 0)
    d.add_heading("Objectives", level=1)
    d.add_paragraph("Improve self-service conversion by 20% for returning customers.", style="List Bullet")
    d.add_paragraph("Reduce support ticket volume related to account access.", style="List Bullet")
    d.add_heading("Timeline", level=1)
    d.add_paragraph("Phase 1: Discovery and requirements - due 15 Nov 2026", style="List Bullet")
    d.add_paragraph("Phase 2: Design and prototyping - due 10 Dec 2026", style="List Bullet")
    d.add_heading("Deliverables", level=1)
    d.add_paragraph("UX research report (Owner: Meera Shah)", style="List Bullet")
    d.add_paragraph("Clickable prototype (Owner: Daniel Lee)", style="List Bullet")
    d.add_heading("Roles and Responsibilities", level=1)
    d.add_paragraph("Sam Patel is responsible for backend API development.")
    t = d.add_table(rows=3, cols=3)
    for r, vals in enumerate([("Name", "Role", "Responsibility"),
                              ("Meera Shah", "UX Lead", "User research and design system"),
                              ("Daniel Lee", "Frontend Engineer", "Prototype and UI implementation")]):
        for c, v in enumerate(vals):
            t.cell(r, c).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf(lines) -> bytes:
    content = "BT /F1 11 Tf 50 780 Td 14 TL " + " ".join(
        "(" + l.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ") Tj T*" for l in lines) + " ET"
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs)+1}\n0000000000 65535 f \n".encode() + b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def texts(items, key):
    return " || ".join(str(i[key]).lower() for i in items)


def check(label, cond):
    print(f"  [{'PASS' if cond else 'FAIL'}] {label}")
    return cond


def run():
    records = (extract_docx(make_docx(), "project_plan.docx") + extract_txt(PLAN_TXT.encode(), "course_plan.txt")
               + extract_csv(CSV_TRACKER.encode(), "tracker.csv") + extract_txt(MEETING_TXT.encode(), "team_sync.txt") + extract_pdf(make_pdf(PDF_LINES), "status_week6.pdf"))
    scope = ScopeDeliverableAgent().analyze(records)
    risk = RiskForecastingAgent().analyze(records, scope["milestones"])
    blk = BlockerActionItemAgent().analyze(records)
    g, m = texts(scope["goals"], "Goal / Objective"), texts(scope["milestones"], "Milestone / Phase")
    d, r = texts(scope["deliverables"], "Deliverable"), texts(scope["responsibilities"], "Role")
    for k, key in [("goals", "Goal / Objective"), ("milestones", "Milestone / Phase"), ("deliverables", "Deliverable"),
                   ("responsibilities", "Role")]:
        print(f"{k}: {len(scope[k])}")
        for i in scope[k]:
            print(f"    - {i[key]}  [{i['Source']} · {i['Location']}]")
    for k, key in [("blockers", "Blocker"), ("decisions", "Decision"), ("action_items", "Action Item")]:
        print(f"{k}: {len(blk[k])}")
        for i in blk[k]:
            print(f"    - {i[key]}  [{i['Source']} · {i['Location']}]")
    ok = all([
        check("DOCX objectives (bullets under heading)", "self-service conversion" in g and "support ticket" in g),
        check("TXT objectives", "grounded conversational assistant" in g),
        check("PDF goal sentence ('The aim of ... is to')", "reporting module to pilot users" in g),
        check("DOCX phases under Timeline", "phase 1" in m and "phase 2" in m),
        check("TXT 'Milestone 3/4 (Week ...)' headings", "milestone 3" in m and "milestone 4" in m),
        check("CSV milestones by Type column", "requirements sign-off" in m and "beta release" in m),
        check("PDF 'Milestone: ... by <date>'", "pilot launch" in m),
        check("DOCX deliverables under heading", "ux research report" in d and "clickable prototype" in d),
        check("TXT milestone tasks become deliverables", "insights and risk summary dashboard" in d and "incremental document upload" in d),
        check("CSV deliverables by Type column", "test plan and test report" in d),
        check("Evaluation criteria NOT treated as deliverables", "accuracy and completeness" not in d),
        check("DOCX roles table", "ux lead" in r and "frontend engineer" in r),
        check("'X is responsible for Y' sentences", "sam patel" in r and "qa lead" in r),
        check("Milestone dates parsed", any(i["Target Date"] not in ("TBD", "") for i in scope["milestones"])),
        check("Deliverable owners parsed", "meera shah" in texts(scope["deliverables"], "Owner")),
        check("Risks detected", len(risk["risks"]) > 0),
        check("Blocker detected in free text", "vendor api outage" in texts(blk["blockers"], "Blocker")),
        check("Blocker from 'unable to' phrasing", "unable to run performance tests" in texts(blk["blockers"], "Blocker")),
        check("CSV row with Status=Blocked is a blocker", "login timeout bug" in texts(blk["blockers"], "Blocker")),
        check("Decision from 'Decisions needed' section", "move the public launch" in texts(blk["decisions"], "Decision")),
        check("Actions from 'Next steps' with owners", "ravi" in texts(blk["action_items"], "Assignee") and "priya" in texts(blk["action_items"], "Assignee")),
        check("Unowned next step still captured", "update the project risk log" in texts(blk["action_items"], "Action Item")),
        check("Action due date parsed", any(a["Due Date"] for a in blk["action_items"] if "login timeout" in a["Action Item"].lower())),
        check("Every item cites a source", all(i.get("Source") for k in scope for i in scope[k])),
    ])
    assert ok, "Generic-document extraction checks failed"
    print("=== GENERIC DOCUMENT EXTRACTION TEST PASSED ===")


if __name__ == "__main__":
    run()

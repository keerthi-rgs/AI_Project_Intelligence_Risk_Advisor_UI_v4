import os
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import List, Dict, Any

import pandas as pd
import streamlit as st
import plotly.express as px
import plotly.graph_objects as go
from pypdf import PdfReader
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from agents import (
    ScopeDeliverableAgent,
    RiskForecastingAgent,
    BlockerActionItemAgent,
    DocumentationGenerationAgent,
    ProjectHealthScorer,
    ConversationalProjectIntelligenceAssistant,
)

# ─────────────────────────────────────────────────────────────
#  Plotly helpers
# ─────────────────────────────────────────────────────────────
def render_risk_matrix_chart(risks: List[Dict[str, Any]], severity_filter: str = "All"):
    """Renders interactive Risk Priority Matrix scatter plot (Probability vs Impact)."""
    if not risks:
        return None
    df = pd.DataFrame(risks)

    if "Probability" not in df.columns:
        df["Probability"] = 0.5
    if "Impact" not in df.columns:
        df["Impact"] = 3.0
    if "Risk ID" not in df.columns:
        df["Risk ID"] = [f"RSK-{i+1:02d}" for i in range(len(df))]
    if "Mitigation Strategy" not in df.columns:
        df["Mitigation Strategy"] = "Monitor sprint velocity."

    if severity_filter != "All":
        df = df[df["Severity"].str.lower() == severity_filter.lower()]

    if df.empty:
        return None

    color_map = {
        "Critical": "#dc2626",
        "High":     "#ef4444",
        "Medium":   "#f59e0b",
        "Low":      "#10b981",
    }

    # map generic "High"→"High", fill missing
    if "Severity" in df.columns:
        df["Severity"] = df["Severity"].fillna("Medium")

    fig = px.scatter(
        df,
        x="Probability",
        y="Impact",
        color="Severity",
        text="Risk ID",
        color_discrete_map=color_map,
        hover_data={
            "Risk Category": True,
            "Risk Description": True,
            "Mitigation Strategy": True,
            "Probability": ":.0%",
            "Impact": ":.1f",
        },
        title="Risk Priority Matrix — Probability vs. Impact",
    )

    fig.update_traces(
        marker=dict(size=18, line=dict(width=2, color="#0f172a")),
        textposition="top center",
        textfont=dict(size=10, color="#0f172a", family="Inter, sans-serif"),
    )

    # Critical Hazard zone
    fig.add_shape(
        type="rect", x0=0.6, y0=3.5, x1=1.02, y1=5.2,
        fillcolor="rgba(239, 68, 68, 0.10)",
        line=dict(color="rgba(239,68,68,0.4)", width=1, dash="dot"),
    )
    fig.add_annotation(
        x=0.81, y=4.92, text="CRITICAL HAZARD ZONE", showarrow=False,
        font=dict(color="#dc2626", size=10, family="Inter, sans-serif"),
    )

    # Low priority zone
    fig.add_shape(
        type="rect", x0=0.0, y0=0.8, x1=0.45, y1=2.5,
        fillcolor="rgba(16,185,129,0.08)",
        line=dict(color="rgba(16,185,129,0.3)", width=1, dash="dot"),
    )
    fig.add_annotation(
        x=0.22, y=1.1, text="LOW PRIORITY ZONE", showarrow=False,
        font=dict(color="#059669", size=10, family="Inter, sans-serif"),
    )

    fig.update_layout(
        xaxis=dict(title="Probability →", range=[0.0, 1.05], tickformat=".0%"),
        yaxis=dict(title="Impact (1.0–5.0) →", range=[0.5, 5.3]),
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="#f8fafc",
        font=dict(family="Inter, sans-serif", color="#0f172a"),
        margin=dict(l=40, r=20, t=50, b=40),
        legend=dict(
            orientation="h", yanchor="bottom", y=1.0,
            xanchor="right", x=1.0,
            title_text="",
        ),
        height=420,
    )
    return fig


def render_risk_category_chart(risks: List[Dict[str, Any]]):
    """Renders Risk Distribution by Category."""
    if not risks:
        return None
    df = pd.DataFrame(risks)
    if "Risk Category" not in df.columns:
        return None
    cat_counts = df["Risk Category"].value_counts().reset_index()
    cat_counts.columns = ["Risk Category", "Count"]

    fig = px.bar(
        cat_counts, x="Count", y="Risk Category",
        orientation="h", color="Count",
        color_continuous_scale=[[0, "#fca5a5"], [1, "#b91c1c"]],
        title="Risk Signals by Category",
    )
    fig.update_layout(
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="#f8fafc",
        font=dict(family="Inter, sans-serif", color="#0f172a"),
        margin=dict(l=40, r=20, t=50, b=40),
        coloraxis_showscale=False,
        height=380,
    )
    return fig


# ─────────────────────────────────────────────────────────────
#  Page config & global CSS
# ─────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="AI Project Intelligence and Risk Advisor",
    page_icon="🧰",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

html, body, [class*="css"] {
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
}

/* ── Main area ── */
.stApp, [data-testid="stHeader"] {
    background-color: #f1f5f9 !important;
}
.block-container {
    padding: 1.2rem 1.6rem 2rem 1.6rem;
    max-width: 100%;
}

/* ── Sidebar ── */
section[data-testid="stSidebar"] {
    background: #ffffff !important;
    border-right: 1px solid #e2e8f0;
    min-width: 220px !important;
    max-width: 240px !important;
}
section[data-testid="stSidebar"] .block-container {
    padding: 0 0 1rem 0 !important;
}

/* ── Sidebar brand ── */
.sb-brand {
    display: flex;
    align-items: center;
    gap: 8px;
    padding: 18px 16px 14px 16px;
    font-size: 15px;
    font-weight: 800;
    color: #0f172a;
    border-bottom: 1px solid #f1f5f9;
}
.sb-brand-sub { font-size: 10px; font-weight: 500; color: #64748b; display:block; margin-top:1px; }

/* ── Sidebar sections ── */
.sb-section {
    font-size: 10px;
    font-weight: 700;
    color: #94a3b8;
    text-transform: uppercase;
    letter-spacing: .08em;
    padding: 14px 16px 4px 16px;
}

/* ── Nav items ── */
.nav-item {
    display: flex;
    align-items: center;
    gap: 10px;
    padding: 9px 16px;
    border-radius: 0;
    font-size: 13.5px;
    font-weight: 500;
    color: #475569;
    cursor: pointer;
    transition: background .15s, color .15s;
    text-decoration: none;
}
.nav-item:hover { background: #f8fafc; color: #0f172a; }
.nav-item.active {
    background: #eff6ff;
    color: #2563eb;
    font-weight: 700;
    border-right: 3px solid #2563eb;
}
.nav-icon { font-size: 16px; width: 20px; text-align: center; }

/* ── Page header bar ── */
.page-header-bar {
    display: flex;
    align-items: center;
    justify-content: space-between;
    background: #ffffff;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
    padding: 14px 20px;
    margin-bottom: 18px;
    box-shadow: 0 1px 4px rgba(0,0,0,0.04);
}
.page-title { font-size: 22px; font-weight: 800; color: #0f172a; margin:0; }
.page-sub   { font-size: 13px; color: #64748b; margin: 3px 0 0 0; }
.header-actions { display: flex; gap: 8px; align-items: center; }

/* ── Badges / pills ── */
.badge { display:inline-block; padding: 4px 13px; border-radius: 20px; font-size: 12px; font-weight: 700; }
.badge-red    { background:#fef2f2; color:#dc2626; border:1px solid #fecaca; }
.badge-orange { background:#fffbeb; color:#d97706; border:1px solid #fde68a; }
.badge-green  { background:#f0fdf4; color:#16a34a; border:1px solid #bbf7d0; }
.badge-blue   { background:#eff6ff; color:#2563eb; border:1px solid #bfdbfe; }

/* ── Action buttons in header ── */
.btn-primary {
    background: #2563eb; color: #fff; border: none;
    padding: 7px 16px; border-radius: 8px; font-size: 13px;
    font-weight: 600; cursor: pointer; font-family: inherit;
}
.btn-outline {
    background: #fff; color: #374151; border: 1px solid #d1d5db;
    padding: 7px 14px; border-radius: 8px; font-size: 13px;
    font-weight: 600; cursor: pointer; font-family: inherit;
}

/* ── Stat cards ── */
.stat-card {
    background: #ffffff;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
    padding: 16px 18px;
    text-align: center;
    box-shadow: 0 1px 3px rgba(0,0,0,0.03);
}
.stat-label { font-size: 12px; font-weight: 600; color: #64748b; text-transform:uppercase; letter-spacing:.04em; }
.stat-value { font-size: 30px; font-weight: 800; color: #0f172a; margin: 4px 0 0 0; line-height:1.1; }
.stat-value.red    { color: #dc2626; }
.stat-value.amber  { color: #d97706; }
.stat-value.green  { color: #16a34a; }
.stat-value.blue   { color: #2563eb; }
.stat-value.purple { color: #7c3aed; }

/* ── Section card ── */
.section-card {
    background: #ffffff;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
    padding: 18px 20px;
    margin-bottom: 16px;
    box-shadow: 0 1px 4px rgba(0,0,0,0.03);
}
.section-card-title {
    font-size: 15px; font-weight: 700; color: #0f172a; margin: 0 0 2px 0;
}
.section-card-sub {
    font-size: 12px; color: #64748b; margin: 0 0 14px 0;
}

/* ── Filter pill buttons ── */
.filter-pills { display: flex; gap: 6px; flex-wrap: wrap; margin-bottom: 12px; }
.pill {
    border: 1px solid #e2e8f0; background: #fff;
    color: #475569; padding: 4px 14px; border-radius: 20px;
    font-size: 12px; font-weight: 600; cursor: pointer;
}
.pill.active { background:#2563eb; color:#fff; border-color:#2563eb; }
.pill.critical { border-color:#fecaca; color:#dc2626; }
.pill.critical.active { background:#dc2626; color:#fff; }
.pill.high { border-color:#fed7aa; color:#d97706; }
.pill.high.active { background:#d97706; color:#fff; }
.pill.moderate { border-color:#fde68a; color:#b45309; }
.pill.moderate.active { background:#b45309; color:#fff; }
.pill.low { border-color:#bbf7d0; color:#16a34a; }
.pill.low.active { background:#16a34a; color:#fff; }

/* ── Metric cards ── */
.metric-card {
    background: #ffffff; border: 1px solid #e2e8f0;
    border-radius: 10px; padding: 14px 16px;
    box-shadow: 0 1px 3px rgba(0,0,0,0.03);
}
.metric-label { color: #64748b; font-size: 11px; font-weight: 600; text-transform:uppercase; letter-spacing:.03em; }
.metric-value { color: #0f172a; font-size: 24px; font-weight: 800; margin-top: 4px; }

/* ── Streamlit overrides ── */
div[data-baseweb="tab-list"] { display: none !important; }
#MainMenu, footer, [data-testid="stToolbar"] { visibility: hidden; }

/* ── Streamlit sidebar radio override to look like nav ── */
div[data-testid="stSidebarContent"] .stRadio > div {
    display: flex; flex-direction: column; gap: 0px;
}
div[data-testid="stSidebarContent"] .stRadio label {
    display: flex; align-items: center; gap: 10px;
    padding: 8px 16px; font-size: 13.5px; font-weight: 500;
    color: #475569; cursor: pointer; border-radius: 0;
    transition: background .15s;
}
div[data-testid="stSidebarContent"] .stRadio label:hover { background: #f8fafc; color: #0f172a; }
div[data-testid="stSidebarContent"] .stRadio [data-baseweb="radio"] { display: none; }

/* upload box */
[data-testid="stFileUploader"] {
    background: #fff; border: 2px dashed #cbd5e1;
    border-radius: 10px; padding: 10px;
}
[data-testid="stFileUploader"]:hover { border-color: #2563eb; }
div[data-testid="stExpander"] {
    border: 1px solid #e2e8f0; border-radius: 10px;
    background: #fff; box-shadow: 0 1px 2px rgba(0,0,0,0.02);
}
</style>
""", unsafe_allow_html=True)

st.markdown("""
<style>
/* ── UX refresh ── */
[data-testid="stHeader"] { background: transparent !important; pointer-events: none; }
[data-testid="stHeader"] button, [data-testid="stHeader"] a { pointer-events: auto; }
.block-container, [data-testid="stMainBlockContainer"] { padding-top: 1.4rem !important; }
.page-header-bar { margin-top: 0; }
.block-container h3 { font-size: 1.1rem !important; font-weight: 700 !important; }
.block-container h4 { font-size: 1rem !important; font-weight: 700 !important; }

/* sidebar */
section[data-testid="stSidebar"] { min-width: 250px !important; max-width: 260px !important; }
[data-testid="stSidebarHeader"] { height: 2.25rem !important; min-height: 0 !important; padding: .4rem .75rem 0 !important; margin: 0 !important; }
[data-testid="stSidebarUserContent"] { padding: 0 .75rem 1rem .75rem !important; }
section[data-testid="stSidebar"] [data-testid="stVerticalBlock"] { gap: 2px !important; }
.sb-brand { padding: 4px 6px 14px 6px; gap: 10px; font-size: 14px; line-height: 1.3; border-bottom: 1px solid #e2e8f0; margin-bottom: 4px; }
.sb-logo { flex: 0 0 34px; height: 34px; border-radius: 9px; background: linear-gradient(135deg,#2563eb,#7c3aed);
           display:flex; align-items:center; justify-content:center; font-size: 18px; color: #fff; }
.sb-section { padding: 16px 8px 8px 8px; }
section[data-testid="stSidebar"] .stMarkdown { margin-bottom: 0 !important; }
section[data-testid="stSidebar"] .stButton { position: relative; z-index: 0; }
section[data-testid="stSidebar"] [data-testid="stElementContainer"]:has(.sb-section) { margin-bottom: 4px; }
section[data-testid="stSidebar"] .stButton button {
    justify-content: flex-start !important; background: transparent !important; border: none !important;
    box-shadow: none !important; color: #475569 !important; padding: 6px 10px !important;
    min-height: 36px !important; border-radius: 8px !important; font-weight: 500 !important;
}
section[data-testid="stSidebar"] .stButton button > div { justify-content: flex-start !important; width: 100%; }
section[data-testid="stSidebar"] .stButton button p { font-size: 13.5px !important; text-align: left; }
section[data-testid="stSidebar"] .stButton button:hover { background: #f1f5f9 !important; color: #0f172a !important; }
section[data-testid="stSidebar"] .stButton button[kind="primary"] {
    background: #eff6ff !important; color: #1d4ed8 !important; box-shadow: inset 3px 0 0 #2563eb !important;
}
section[data-testid="stSidebar"] .stButton button[kind="primary"] p { font-weight: 700 !important; }
.kb-card { background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 10px; padding: 12px 14px; margin-top: 6px; }
.kb-title { display:flex; align-items:center; justify-content:space-between; font-size: 11px; font-weight: 700;
            color: #64748b; text-transform: uppercase; letter-spacing: .06em; margin-bottom: 8px; }
.kb-row { display:flex; justify-content:space-between; font-size: 13px; color:#475569; padding: 2px 0; }
.kb-row b { color:#0f172a; }
.dot { display:inline-block; width:8px; height:8px; border-radius:50%; margin-right:5px; }
.sb-footer { font-size: 11px; color: #94a3b8; padding: 12px 6px 0 6px; line-height: 1.4; }

/* cards (st.container(border=True, key="card_…")) */
div[class*="st-key-card"] {
    background: #ffffff !important; border: 1px solid #e2e8f0 !important; border-radius: 12px !important;
    box-shadow: 0 1px 3px rgba(15,23,42,.04); padding: 16px 18px !important;
}
.section-card-title { font-size: 15px; font-weight: 700; color:#0f172a; }
.section-card-sub { margin: 2px 0 10px 0; }

/* metrics */
[data-testid="stMetric"] { background:#f8fafc; border:1px solid #eef2f7; border-radius:10px; padding:10px 14px; }
[data-testid="stMetricLabel"] p { font-size:12px !important; font-weight:600 !important; color:#64748b !important;
                                  text-transform:uppercase; letter-spacing:.04em; }
[data-testid="stMetricValue"] { font-size: 1.6rem !important; font-weight: 800 !important; color:#0f172a; }

/* buttons */
.stButton button, .stDownloadButton button { border-radius: 8px !important; font-weight: 600 !important; }

/* empty state */
.empty-state { text-align:center; padding: 28px 10px 12px 10px; }
.empty-icon { width:64px; height:64px; margin:0 auto 12px auto; border-radius:16px; background:#eff6ff;
              display:flex; align-items:center; justify-content:center; font-size:30px; }
.empty-title { font-size: 20px; font-weight: 800; color:#0f172a; margin-bottom: 6px; }
.empty-msg { font-size: 14px; color:#64748b; max-width: 560px; margin: 0 auto 8px auto; line-height: 1.5; }
.empty-steps { display:flex; gap:10px; justify-content:center; flex-wrap:wrap; margin: 14px 0 6px 0; }
.empty-step { font-size:12px; color:#475569; background:#f8fafc; border:1px solid #e2e8f0; border-radius:999px; padding:4px 12px; }

/* score bars */
.score-row { margin: 4px 0 14px 0; }
.score-head { display:flex; justify-content:space-between; font-size: 13.5px; font-weight: 600; color:#334155; margin-bottom: 6px; }
.score-track { height: 10px; border-radius: 999px; background:#e2e8f0; overflow:hidden; }
.score-fill { height: 100%; border-radius: 999px; }
.score-hint { font-size: 12px; color:#64748b; margin-top: 4px; }

/* chips */
.chip { display:inline-block; padding:3px 10px; border-radius:999px; background:#eff6ff; color:#1d4ed8;
        font-size:12px; font-weight:600; margin:0 6px 6px 0; border:1px solid #bfdbfe; }
</style>
""", unsafe_allow_html=True)


# ─────────────────────────────────────────────────────────────
#  Data structures / state
# ─────────────────────────────────────────────────────────────
@dataclass
class Chunk:
    chunk_id: str
    source: str
    location: str
    text: str


class FileOnDisk:
    def __init__(self, filepath: str):
        self.filepath = filepath
        self.name = Path(filepath).name

    def getvalue(self) -> bytes:
        with open(self.filepath, "rb") as f:
            return f.read()


# ── Session state init ──
for key, default in [
    ("documents", []),
    ("chunks", []),
    ("vectorizer", None),
    ("matrix", None),
    ("last_results", []),
    ("ingestion_log", []),
    ("chat_history", []),
    ("active_page", "Dashboard"),
    ("risk_filter", "All"),
    ("indexed_sources", []),
    ("kb_version", 0),
    ("kb_history", []),
]:
    if key not in st.session_state:
        st.session_state[key] = default

# ── Agent init ──
if "scope_agent"             not in st.session_state: st.session_state.scope_agent             = ScopeDeliverableAgent()
if "risk_agent"              not in st.session_state: st.session_state.risk_agent              = RiskForecastingAgent()
if "blocker_agent"           not in st.session_state: st.session_state.blocker_agent           = BlockerActionItemAgent()
if "doc_gen_agent"           not in st.session_state: st.session_state.doc_gen_agent           = DocumentationGenerationAgent()
if "health_scorer"           not in st.session_state: st.session_state.health_scorer           = ProjectHealthScorer()
if "conversational_assistant" not in st.session_state: st.session_state.conversational_assistant = ConversationalProjectIntelligenceAssistant()


# ─────────────────────────────────────────────────────────────
#  Sidebar navigation
# ─────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("""
    <div class="sb-brand">
        <div class="sb-logo">🧰</div>
        <div>AI Project Intelligence and Risk Advisor</div>
    </div>
    """, unsafe_allow_html=True)

    def nav_item(label: str, icon: str, page: str):
        active = "active" if st.session_state.active_page == page else ""
        if st.button(f"{icon}  {label}", key=f"nav_{page}",
                     width="stretch",
                     type="primary" if active else "secondary"):
            st.session_state.active_page = page
            st.rerun()

    # ── MAIN ──
    st.markdown('<div class="sb-section">MAIN</div>', unsafe_allow_html=True)
    nav_item("Dashboard",          "📊", "Dashboard")
    nav_item("Documents",          "📁", "Documents")
    nav_item("Project Intelligence","🧠", "ProjectIntelligence")

    # ── ANALYZE ──
    st.markdown('<div class="sb-section">ANALYZE</div>', unsafe_allow_html=True)
    nav_item("Risks",              "⚠️",  "Risks")
    nav_item("Blockers & Action Items", "🚧", "Blockers")
    nav_item("Project Health",     "🩺", "Health")

    # ── AI & AGENTS ──
    st.markdown('<div class="sb-section">AI & AGENTS</div>', unsafe_allow_html=True)
    nav_item("AI Assistant",       "💬", "AIAssistant")
    nav_item("Reports & Doc Gen.", "📄", "DocGen")

    # ── SYSTEM & SPEC ──
    st.markdown('<div class="sb-section">SYSTEM & SPEC</div>', unsafe_allow_html=True)
    nav_item("Architecture & Milestones", "🗺️", "Architecture")
    nav_item("Settings",           "⚙️", "Settings")

    _n_files = len({c.source for c in st.session_state.chunks})
    _ready = bool(st.session_state.chunks)
    st.markdown(f"""
    <div class="kb-card">
        <div class="kb-title"><span>Knowledge base</span>
            <span style="text-transform:none; letter-spacing:0; font-weight:600; color:{'#16a34a' if _ready else '#94a3b8'};">
                <span class="dot" style="background:{'#16a34a' if _ready else '#cbd5e1'};"></span>{'Ready' if _ready else 'Empty'}
            </span>
        </div>
        <div class="kb-row"><span>Source files</span><b>{_n_files}</b></div>
        <div class="kb-row"><span>Indexed chunks</span><b>{len(st.session_state.chunks)}</b></div>
        <div class="kb-row"><span>Version</span><b>v{st.session_state.kb_version}</b></div>
    </div>
    <div class="sb-footer">AI Project Intelligence and Risk Advisor · Document intelligence &amp; risk forecasting</div>
    """, unsafe_allow_html=True)


# ─────────────────────────────────────────────────────────────
#  Extraction helpers
# ─────────────────────────────────────────────────────────────
def normalize(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def extract_pdf(data: bytes, filename: str):
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = normalize(page.extract_text() or "")
        if text:
            pages.append({"source": filename, "location": f"Page {i}", "text": text})
    return pages

def extract_docx(data: bytes, filename: str):
    doc = Document(io.BytesIO(data))
    records = []
    for i, p in enumerate(doc.paragraphs, start=1):
        text = normalize(p.text)
        if text:
            records.append({"source": filename, "location": f"Paragraph {i}", "text": text})
    for ti, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [normalize(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        table_text = normalize("\n".join(rows))
        if table_text:
            records.append({"source": filename, "location": f"Table {ti}", "text": table_text})
    if not records:
        import zipfile as _zipfile
        from xml.etree import ElementTree as ET
        try:
            with _zipfile.ZipFile(io.BytesIO(data)) as z:
                xml = z.read("word/document.xml")
            root_xml = ET.fromstring(xml)
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            texts = [t.text for t in root_xml.findall(".//w:t", ns) if t.text]
            fallback = normalize(" ".join(texts))
            if fallback:
                records.append({"source": filename, "location": "DOCX XML content", "text": fallback})
        except Exception:
            pass
    return records

def extract_txt(data: bytes, filename: str):
    text = normalize(data.decode("utf-8", errors="ignore"))
    return [{"source": filename, "location": "Text file", "text": text}] if text else []

def extract_csv(data: bytes, filename: str):
    df = pd.read_csv(io.BytesIO(data))
    records = []
    for idx, row in df.fillna("").iterrows():
        parts = [f"{col}: {row[col]}" for col in df.columns if str(row[col]).strip()]
        if parts:
            records.append({"source": filename, "location": f"CSV Row {idx+2}", "text": normalize(" | ".join(parts))})
    return records

def extract_file(uploaded_file):
    name = uploaded_file.name
    data = uploaded_file.getvalue()
    ext = Path(name).suffix.lower()
    if ext == ".pdf":  return extract_pdf(data, name)
    if ext == ".docx": return extract_docx(data, name)
    if ext == ".txt":  return extract_txt(data, name)
    if ext == ".csv":  return extract_csv(data, name)
    raise ValueError(f"Unsupported format: {ext}")


# ─────────────────────────────────────────────────────────────
#  Chunking & Indexing
# ─────────────────────────────────────────────────────────────
CHUNK_SIZE    = 500
CHUNK_OVERLAP = 100
TOP_K         = 4

def chunk_text(text: str, size: int, overlap: int):
    text = normalize(text)
    if not text:
        return []
    step = max(1, size - overlap)
    result, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        piece = text[start:end].strip()
        if end < len(text):
            candidates = [piece.rfind(". "), piece.rfind("\n"), piece.rfind(", ")]
            best = max(c for c in candidates if c > size // 2) if any(c > size // 2 for c in candidates) else -1
            if best > 0:
                piece = piece[:best + 1].strip()
                end = start + best + 1
        if piece:
            result.append(piece)
        start = end if end >= len(text) else start + step
    return result

def build_chunks(records, size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    chunks = []
    for record in records:
        pieces = chunk_text(record["text"], size, overlap)
        for i, piece in enumerate(pieces, start=1):
            chunks.append(Chunk(
                chunk_id=f"{Path(record['source']).stem}_C{i:03d}",
                source=record["source"],
                location=record["location"],
                text=piece,
            ))
    return chunks

def rebuild_index(chunks):
    if not chunks:
        st.session_state.vectorizer = None
        st.session_state.matrix = None
        return None, None
    texts = [c.text for c in chunks]
    vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1, 2))
    matrix = vectorizer.fit_transform(texts)
    st.session_state.vectorizer = vectorizer
    st.session_state.matrix     = matrix
    return vectorizer, matrix


def merge_documents_incremental(existing, incoming):
    """Replace units from the same source file, then append new sources."""
    incoming_sources = {r.get("source") for r in incoming}
    kept = [r for r in existing if r.get("source") not in incoming_sources]
    return kept + incoming


def knowledge_base_fingerprint(documents):
    return hash(tuple((d.get("source", ""), d.get("location", ""), d.get("text", "")) for d in documents))


def apply_ingestion(all_records, log, incremental=True):
    """Merge or replace the knowledge base, then rebuild the RAG index automatically."""
    if incremental and st.session_state.documents:
        before_sources = {r.get("source") for r in st.session_state.documents}
        incoming_sources = {r.get("source") for r in all_records}
        replaced = sorted(before_sources & incoming_sources)
        added = sorted(incoming_sources - before_sources)
        st.session_state.documents = merge_documents_incremental(st.session_state.documents, all_records)
        st.session_state.ingestion_log.extend(log)
        action = (
            f"Knowledge base updated automatically. "
            f"Added {len(added)} file(s), replaced {len(replaced)} file(s), "
            f"ingested {len(all_records)} content units."
        )
        history_action = "incremental"
    else:
        st.session_state.documents = list(all_records)
        st.session_state.ingestion_log = list(log)
        action = f"Knowledge base replaced with {len(all_records)} units from {len(log)} file(s)."
        history_action = "overwrite"
        replaced, added = [], sorted({r.get("source") for r in all_records})

    st.session_state.chunks = build_chunks(st.session_state.documents)
    rebuild_index(st.session_state.chunks)
    st.session_state.indexed_sources = sorted({c.source for c in st.session_state.chunks})
    st.session_state.kb_version = int(st.session_state.kb_version) + 1
    st.session_state.kb_history.append({
        "Version": st.session_state.kb_version,
        "Mode": history_action,
        "Added": ", ".join(added) if added else "—",
        "Replaced": ", ".join(replaced) if replaced else "—",
        "Units": len(st.session_state.documents),
        "Chunks": len(st.session_state.chunks),
    })
    get_all_analysis.clear()
    return action

def retrieve(query, k=TOP_K):
    if not st.session_state.chunks or st.session_state.vectorizer is None or st.session_state.matrix is None:
        return []
    qvec = st.session_state.vectorizer.transform([query])
    scores = cosine_similarity(qvec, st.session_state.matrix)[0]
    order = scores.argsort()[::-1][:k]
    return [{"chunk": st.session_state.chunks[idx], "score": float(scores[idx])} for idx in order]


# ─────────────────────────────────────────────────────────────
#  Helper: load sample data
# ─────────────────────────────────────────────────────────────
def load_sample_data():
    sample_dir = Path("sample_data")
    if not sample_dir.exists():
        return False, "sample_data directory not found."
    files = list(sample_dir.glob("*.*"))
    all_records, log = [], []
    for fpath in files:
        try:
            mock = FileOnDisk(str(fpath))
            recs = extract_file(mock)
            all_records.extend(recs)
            log.append({"File": mock.name, "Status": "Success", "Units": len(recs)})
        except Exception as e:
            log.append({"File": fpath.name, "Status": f"Failed: {e}", "Units": 0})
    apply_ingestion(all_records, log, incremental=False)
    return True, f"Loaded {len(files)} files → {len(st.session_state.documents)} units, {len(st.session_state.chunks)} chunks."


def load_progress_update():
    update_path = Path("sample_data") / "progress_update.txt"
    if not update_path.exists():
        return False, "sample_data/progress_update.txt not found."
    mock = FileOnDisk(str(update_path))
    recs = extract_file(mock)
    log = [{"File": mock.name, "Status": "Success", "Units": len(recs)}]
    msg = apply_ingestion(recs, log, incremental=True)
    return True, msg


# ─────────────────────────────────────────────────────────────
#  Helper: get live agent data
# ─────────────────────────────────────────────────────────────
@st.cache_data(show_spinner=False)
def get_all_analysis(docs_hash: int, kb_version: int):
    docs = st.session_state.documents
    scope_data   = st.session_state.scope_agent.analyze(docs)
    risk_data    = st.session_state.risk_agent.analyze(docs, scope_data["milestones"])
    blocker_data = st.session_state.blocker_agent.analyze(docs)
    health_data  = st.session_state.health_scorer.calculate_health_score(docs, scope_data, risk_data, blocker_data)
    return scope_data, risk_data, blocker_data, health_data



# ─────────────────────────────────────────────────────────────
#  UI helpers
# ─────────────────────────────────────────────────────────────
def page_header(icon: str, title: str, sub: str, right_html: str = ""):
    st.markdown(f"""
    <div class="page-header-bar">
        <div>
            <div class="page-title">{icon} {title}</div>
            <div class="page-sub">{sub}</div>
        </div>
        <div class="header-actions">{right_html}</div>
    </div>
    """, unsafe_allow_html=True)


def section_title(title: str, sub: str = ""):
    html = f'<div class="section-card-title">{title}</div>'
    if sub:
        html += f'<div class="section-card-sub">{sub}</div>'
    st.markdown(html, unsafe_allow_html=True)


def go_to(page_key: str):
    st.session_state.active_page = page_key
    st.rerun()


def empty_state(icon: str, title: str, message: str, key: str):
    with st.container(border=True, key=f"card_empty_{key}"):
        st.markdown(f"""
        <div class="empty-state">
            <div class="empty-icon">{icon}</div>
            <div class="empty-title">{title}</div>
            <div class="empty-msg">{message}</div>
            <div class="empty-steps">
                <span class="empty-step">1 · Upload PDF, DOCX, TXT or CSV</span>
                <span class="empty-step">2 · Agents extract scope, risks &amp; actions</span>
                <span class="empty-step">3 · Explore insights &amp; ask questions</span>
            </div>
        </div>
        """, unsafe_allow_html=True)
        _, b1, b2, _ = st.columns([1, 2, 2, 1])
        with b1:
            if st.button("Upload my documents", icon="📁", type="primary",
                         width="stretch", key=f"empty_upload_{key}"):
                go_to("Documents")
        with b2:
            if st.button("Try the sample dataset", icon="🚀",
                         width="stretch", key=f"empty_sample_{key}"):
                ok, msg = load_sample_data()
                if ok:
                    st.rerun()
                else:
                    st.error(msg)


def score_color(value: float) -> str:
    return "#dc2626" if value < 50 else "#d97706" if value < 75 else "#16a34a"


def score_bar(label: str, value: float, hint: str = ""):
    color = score_color(value)
    width = max(0, min(100, value))
    hint_html = f'<div class="score-hint">{hint}</div>' if hint else ""
    st.markdown(f"""
    <div class="score-row">
        <div class="score-head"><span>{label}</span><span style="color:{color};">{value}/100</span></div>
        <div class="score-track"><div class="score-fill" style="width:{width}%; background:{color};"></div></div>
        {hint_html}
    </div>
    """, unsafe_allow_html=True)


def health_gauge(score: float, height: int = 220, number_size: int = 34):
    fig = go.Figure(go.Indicator(
        mode="gauge+number",
        value=score,
        number={"suffix": "/100", "font": {"size": number_size, "color": "#0f172a"}},
        gauge={
            "axis": {"range": [0, 100], "tickwidth": 1, "tickcolor": "#cbd5e1"},
            "bar": {"color": score_color(score), "thickness": 0.3},
            "bgcolor": "white",
            "borderwidth": 0,
            "steps": [
                {"range": [0, 50], "color": "#fef2f2"},
                {"range": [50, 75], "color": "#fffbeb"},
                {"range": [75, 100], "color": "#f0fdf4"},
            ],
        },
    ))
    fig.update_layout(height=height, margin=dict(l=34, r=34, t=24, b=8),
                      paper_bgcolor="rgba(0,0,0,0)", font=dict(family="Inter, sans-serif"))
    return fig


def show_table(rows, empty_msg: str, key: str = "", download_name: str = "", success_when_empty: bool = False):
    if rows:
        df = pd.DataFrame(rows)
        st.dataframe(df, width="stretch", hide_index=True)
        if download_name:
            st.download_button("Download CSV", df.to_csv(index=False), f"{download_name}.csv", "text/csv",
                               icon="⬇️", key=f"dl_{key or download_name}")
    elif success_when_empty:
        st.success(empty_msg)
    else:
        st.info(empty_msg)


# ─────────────────────────────────────────────────────────────
#  PAGE RENDERER
# ─────────────────────────────────────────────────────────────
page = st.session_state.active_page


# ══════════════════════════════════════════════════════════════
#  DASHBOARD PAGE
# ══════════════════════════════════════════════════════════════
if page == "Dashboard":
    page_header("📊", "Project Insights &amp; Risk Summary",
                "Health score, extracted risks, action items, and scope summary — refreshed automatically when documents are added.")

    if not st.session_state.documents:
        empty_state("📊", "No project documents yet",
                    "Upload project plans, meeting notes, status reports or trackers, or try the built-in sample "
                    "dataset to see the insights dashboard in action.", "dash")
    else:
        docs_hash = knowledge_base_fingerprint(st.session_state.documents)
        scope_data, risk_data, blocker_data, health_data = get_all_analysis(
            docs_hash, int(st.session_state.kb_version)
        )
        forecast = risk_data["forecast"]
        dim = health_data["dimension_scores"]

        badge_cls = health_data["badge_class"]
        f_badge = "badge-red" if forecast["health_status"] == "CRITICAL RISK" else (
                  "badge-orange" if forecast["health_status"] == "AT RISK" else "badge-green")

        k1, k2, k3, k4, k5 = st.columns(5)
        with k1:
            st.markdown(f"""
            <div class="stat-card">
                <div class="stat-label">Health Score</div>
                <div class="stat-value blue">{health_data['overall_score']}<span style="font-size:14px; color:#64748b;">/100</span></div>
                <div style="margin-top:6px;"><span class="badge {badge_cls}">{health_data['health_status']}</span></div>
            </div>""", unsafe_allow_html=True)
        with k2:
            st.markdown(f"""
            <div class="stat-card">
                <div class="stat-label">Delivery Forecast</div>
                <div class="stat-value" style="font-size:17px; color:#2563eb;">{forecast['estimated_slippage']}</div>
                <div style="margin-top:6px;"><span class="badge {f_badge}">{forecast['health_status']}</span></div>
            </div>""", unsafe_allow_html=True)
        with k3:
            st.markdown(f"""
            <div class="stat-card">
                <div class="stat-label">High Severity Risks</div>
                <div class="stat-value red">{forecast['high_risk_count']}</div>
                <div style="font-size:11px; color:#64748b; margin-top:4px;">{len(risk_data['risks'])} total signals</div>
            </div>""", unsafe_allow_html=True)
        with k4:
            st.markdown(f"""
            <div class="stat-card">
                <div class="stat-label">Active Blockers</div>
                <div class="stat-value amber">{blocker_data['summary']['total_blockers']}</div>
                <div style="font-size:11px; color:#64748b; margin-top:4px;">{blocker_data['summary']['total_decisions']} pending decisions</div>
            </div>""", unsafe_allow_html=True)
        with k5:
            st.markdown(f"""
            <div class="stat-card">
                <div class="stat-label">Open Action Items</div>
                <div class="stat-value purple">{blocker_data['summary']['open_action_items']}</div>
                <div style="font-size:11px; color:#64748b; margin-top:4px;">{blocker_data['summary']['total_action_items']} total tracked</div>
            </div>""", unsafe_allow_html=True)

        st.markdown("<div style='height:16px'></div>", unsafe_allow_html=True)

        sc1, sc2 = st.columns(2)
        with sc1:
            with st.container(border=True, key="card_dash_scope"):
                section_title("🧭 Scope Summary", "Goals, milestones, deliverables, and RACI roles found in your documents.")
                g1, g2, g3, g4 = st.columns(4)
                g1.metric("Goals", len(scope_data["goals"]))
                g2.metric("Milestones", len(scope_data["milestones"]))
                g3.metric("Deliverables", len(scope_data["deliverables"]))
                g4.metric("RACI Roles", len(scope_data["responsibilities"]))
                if scope_data["goals"]:
                    st.dataframe(
                        pd.DataFrame(scope_data["goals"])[["Goal / Objective", "Category"]].head(5),
                        width="stretch", hide_index=True,
                    )
                if scope_data["milestones"]:
                    st.caption("Milestones")
                    st.dataframe(
                        pd.DataFrame(scope_data["milestones"])[["Milestone / Phase", "Target Date", "Status"]].head(5),
                        width="stretch", hide_index=True,
                    )
                if st.button("View full scope →", type="tertiary", key="dash_go_pi"):
                    go_to("ProjectIntelligence")
        with sc2:
            with st.container(border=True, key="card_dash_health"):
                section_title("🩺 Health Score Breakdown",
                              "Each dimension is scored 0–100 (higher is healthier) and weighted into the overall score.")
                gcol, bcol = st.columns([2, 3])
                with gcol:
                    st.plotly_chart(health_gauge(health_data["overall_score"], 210, number_size=24), width="stretch",
                                    key="dash_gauge", config={"displayModeBar": False})
                with bcol:
                    score_bar("Scope clarity", dim["scope_clarity"], "30% weight")
                    score_bar("Timeline", dim["timeline_risk"], "40% weight")
                    score_bar("Blockers", dim["blocker_impact"], "30% weight")
                st.markdown(
                    f"<span class='badge {f_badge}'>{forecast['health_status']}</span> &nbsp; "
                    f"<span style='font-size:13px; color:#475569;'>Estimated slippage: <b>{forecast['estimated_slippage']}</b></span>",
                    unsafe_allow_html=True,
                )
                st.caption(health_data["summary_description"])
                if forecast["milestone_forecasts"]:
                    st.dataframe(pd.DataFrame(forecast["milestone_forecasts"]).head(4), width="stretch", hide_index=True)
                if st.button("View health details →", type="tertiary", key="dash_go_health"):
                    go_to("Health")

        r1, r2 = st.columns(2)
        with r1:
            with st.container(border=True, key="card_dash_risks"):
                section_title("⚠️ Extracted Risks", "Top risk signals ranked by severity, with mitigation.")
                if risk_data["risks"]:
                    rdf = pd.DataFrame(risk_data["risks"])
                    cols = [c for c in ["Risk ID", "Severity", "Risk Category", "Risk Description", "Mitigation Strategy"] if c in rdf.columns]
                    severity_rank = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3}
                    if "Severity" in rdf.columns:
                        rdf["_rank"] = rdf["Severity"].map(lambda s: severity_rank.get(str(s), 9))
                        rdf = rdf.sort_values("_rank")
                    st.dataframe(rdf[cols].head(8), width="stretch", hide_index=True)
                else:
                    st.success("No risks detected.")
                if st.button("Open risk analysis →", type="tertiary", key="dash_go_risks"):
                    go_to("Risks")
        with r2:
            with st.container(border=True, key="card_dash_actions"):
                section_title("✅ Action Items", "Assigned follow-ups extracted from meeting notes and backlogs.")
                if blocker_data["action_items"]:
                    adf = pd.DataFrame(blocker_data["action_items"])
                    cols = [c for c in ["Action Item", "Assignee", "Priority", "Status", "Due Date"] if c in adf.columns]
                    st.dataframe(adf[cols].head(8), width="stretch", hide_index=True)
                else:
                    st.info("No action items detected.")
                if st.button("Open blockers & actions →", type="tertiary", key="dash_go_blockers"):
                    go_to("Blockers")

        rc1, rc2 = st.columns([3, 2])
        with rc1:
            with st.container(border=True, key="card_dash_matrix"):
                fig_m = render_risk_matrix_chart(risk_data["risks"])
                if fig_m:
                    st.plotly_chart(fig_m, width="stretch", key="dash_matrix")
                else:
                    st.info("No risk data for matrix.")
        with rc2:
            with st.container(border=True, key="card_dash_cat"):
                fig_c = render_risk_category_chart(risk_data["risks"])
                if fig_c:
                    st.plotly_chart(fig_c, width="stretch", key="dash_cat")
                else:
                    st.info("No risk categories.")

        st.caption(
            f"Knowledge base v{st.session_state.kb_version} · "
            f"{len(st.session_state.indexed_sources)} source file(s) · "
            f"{len(st.session_state.chunks)} indexed chunks. "
            "Upload a new meeting note on the Documents page to refresh this dashboard automatically."
        )


# ══════════════════════════════════════════════════════════════
#  DOCUMENTS PAGE
# ══════════════════════════════════════════════════════════════
elif page == "Documents":
    page_header("📁", "Documents",
                "Add plans, meeting notes or progress updates — the knowledge base and search index update automatically.")

    if "last_ingest_msg" in st.session_state:
        t, m = st.session_state.pop("last_ingest_msg")
        (st.error if t == "error" else st.warning if t == "warning" else st.success)(m)

    up_col, side_col = st.columns([3, 2])
    with up_col:
        with st.container(border=True, key="card_docs_upload"):
            section_title("⬆️ Upload documents",
                          "PDF, DOCX, TXT or CSV. Re-uploading a file with the same name replaces only that file.")
            uploaded = st.file_uploader("Choose project files to ingest", type=["pdf", "docx", "txt", "csv"],
                                        accept_multiple_files=True, label_visibility="collapsed")
            ingest_mode = st.radio(
                "When adding files",
                ["Add to existing knowledge base", "Replace the whole knowledge base"],
                index=0, horizontal=True,
            )
            ingest = st.button(
                f"Ingest {len(uploaded)} file(s)" if uploaded else "Ingest files",
                icon="🔗", type="primary", width="stretch", disabled=not uploaded,
            )

    with side_col:
        with st.container(border=True, key="card_docs_kb"):
            section_title("📚 Knowledge base", "What the agents and assistant are currently reading.")
            k1, k2, k3 = st.columns(3)
            k1.metric("Files", len(st.session_state.indexed_sources))
            k2.metric("Units", len(st.session_state.documents))
            k3.metric("Chunks", len(st.session_state.chunks))
            if st.session_state.indexed_sources:
                st.markdown("".join(f'<span class="chip">📄 {s}</span>' for s in st.session_state.indexed_sources),
                            unsafe_allow_html=True)
            else:
                st.caption("Nothing indexed yet — upload files or load the sample dataset.")
        with st.container(border=True, key="card_docs_sample"):
            section_title("🧪 Sample data", "Five sample files covering every supported format.")
            s1, s2 = st.columns(2)
            with s1:
                if st.button("Load sample dataset", icon="🚀", width="stretch"):
                    ok, msg = load_sample_data()
                    st.session_state.last_ingest_msg = ("success" if ok else "error", msg)
                    st.rerun()
            with s2:
                if st.button("Add progress update", icon="➕", width="stretch",
                             help="Incrementally adds sample_data/progress_update.txt and opens the dashboard."):
                    ok, msg = load_progress_update()
                    if ok:
                        st.session_state.active_page = "Dashboard"
                    else:
                        st.session_state.last_ingest_msg = ("error", msg)
                    st.rerun()

    if ingest and uploaded:
        all_records, log = [], []
        with st.spinner("Extracting text and rebuilding the index…"):
            for file in uploaded:
                try:
                    recs = extract_file(file)
                    all_records.extend(recs)
                    log.append({"File": file.name, "Status": "Success", "Units": len(recs)})
                except Exception as exc:
                    log.append({"File": file.name, "Status": f"Failed: {exc}", "Units": 0})
            incremental = ingest_mode.startswith("Add")
            action = apply_ingestion(all_records, log, incremental=incremental)
        failed = [x for x in log if x["Status"] != "Success"]
        if failed:
            st.session_state.last_ingest_msg = ("error", "Some files failed — see the ingestion log below.")
        elif not all_records:
            st.session_state.last_ingest_msg = ("warning", "No text could be extracted from the selected files.")
        else:
            st.session_state.last_ingest_msg = (
                "success",
                f"{action} Total: {len(st.session_state.documents)} units, {len(st.session_state.chunks)} chunks.",
            )
        st.rerun()

    if st.session_state.kb_history:
        with st.expander(f"🔁 Knowledge base update history ({len(st.session_state.kb_history)})"):
            st.dataframe(pd.DataFrame(st.session_state.kb_history), width="stretch", hide_index=True)
    if st.session_state.ingestion_log:
        with st.expander(f"📊 Ingestion log ({len(st.session_state.ingestion_log)} file events)"):
            st.dataframe(pd.DataFrame(st.session_state.ingestion_log), width="stretch", hide_index=True)
    if st.session_state.documents:
        with st.expander(f"👁️ Content preview (first 15 of {len(st.session_state.documents)} units)"):
            for i, rec in enumerate(st.session_state.documents[:15], 1):
                st.markdown(f"**{i}. {rec['source']} — {rec['location']}**")
                st.caption(rec["text"][:600] + ("…" if len(rec["text"]) > 600 else ""))


# ══════════════════════════════════════════════════════════════
#  PROJECT INTELLIGENCE PAGE
# ══════════════════════════════════════════════════════════════
elif page == "ProjectIntelligence":
    page_header("🧠", "Project Intelligence",
                "Scope, deliverables, milestones, and the RACI responsibility matrix extracted from your documents.")

    if not st.session_state.documents:
        empty_state("🧠", "No scope extracted yet",
                    "Add project documents to extract goals, milestones, deliverables and responsibilities.", "pi")
    else:
        scope_data = st.session_state.scope_agent.analyze(st.session_state.documents)
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Goals & Objectives",   len(scope_data["goals"]))
        m2.metric("Milestones",           len(scope_data["milestones"]))
        m3.metric("Key Deliverables",     len(scope_data["deliverables"]))
        m4.metric("RACI Responsibilities", len(scope_data["responsibilities"]))
        for key, title, sub, empty in [
            ("goals", "📌 Project Goals & Strategic Objectives", "What the project is trying to achieve.", "No goals detected."),
            ("milestones", "📅 Milestones & Delivery Schedule", "Phases, target dates and status.", "No milestones detected."),
            ("deliverables", "📦 Key Deliverables & Owners", "Outputs the project must produce, with owners where stated.", "No deliverables detected."),
            ("responsibilities", "👥 Stakeholder RACI Matrix", "Who is responsible for what.", "No RACI responsibilities detected."),
        ]:
            with st.container(border=True, key=f"card_pi_{key}"):
                section_title(title, sub)
                show_table(scope_data[key], empty, key=f"pi_{key}", download_name=f"scope_{key}")


# ══════════════════════════════════════════════════════════════
#  RISKS PAGE  ← matches the reference screenshot exactly
# ══════════════════════════════════════════════════════════════
elif page == "Risks":

    if not st.session_state.documents:
        page_header("⚠️", "Risks", "Automated risk detection, impact estimation, ranking, and mitigation strategies.")
        empty_state("⚠️", "No risks analysed yet",
                    "Add project documents to detect schedule, resource, dependency, technical and budget risks.", "risks")
    else:
        scope_data   = st.session_state.scope_agent.analyze(st.session_state.documents)
        risk_data    = st.session_state.risk_agent.analyze(st.session_state.documents, scope_data["milestones"])
        blocker_data = st.session_state.blocker_agent.analyze(st.session_state.documents)
        health_data  = st.session_state.health_scorer.calculate_health_score(
            st.session_state.documents, scope_data, risk_data, blocker_data)
        forecast     = risk_data["forecast"]
        risks        = risk_data["risks"]

        h_status = forecast["health_status"]
        h_badge  = "badge-red" if h_status == "CRITICAL RISK" else ("badge-orange" if h_status == "AT RISK" else "badge-green")

        # ── Page header bar with health badge + action buttons ──
        st.markdown(f"""
        <div class="page-header-bar">
            <div>
                <div class="page-title">⚠️ Risks</div>
                <div class="page-sub">Automated vulnerability detection, impact estimation, risk ranking, and assigned mitigation strategies. &nbsp;|&nbsp; Target: <b>{forecast.get('target_date','2026-10-06')}</b></div>
            </div>
            <div class="header-actions">
                <span style="font-size:13px; color:#64748b; font-weight:600;">Project Health</span>
                <span class="badge badge-blue">📋 PRD</span>
                <span class="badge {h_badge}">{'🔴' if 'CRITICAL' in h_status else '🟡' if 'AT RISK' in h_status else '🟢'} {h_status}</span>
            </div>
        </div>
        """, unsafe_allow_html=True)

        # ── Action buttons row ──
        ab1, ab2, ab3, ab_sp = st.columns([2, 2, 2, 4])
        with ab1:
            if st.button("▶ Run Multi-Agent Pipeline", type="primary", width="stretch", key="run_pipeline"):
                st.rerun()
        with ab2:
            if st.button("📄 Document Generation", width="stretch", key="doc_gen_btn"):
                st.session_state.active_page = "DocGen"
                st.rerun()
        with ab3:
            if st.button("💬 Ask Advisor", width="stretch", key="ask_advisor_btn"):
                st.session_state.active_page = "AIAssistant"
                st.rerun()

        st.markdown("<div style='height:6px'></div>", unsafe_allow_html=True)

        # ── Risk severity stat cards ──
        df_risks = pd.DataFrame(risks) if risks else pd.DataFrame()

        def count_sev(label):
            if df_risks.empty or "Severity" not in df_risks.columns:
                return 0
            return int((df_risks["Severity"].str.lower() == label.lower()).sum())

        total_risks   = len(risks)
        critical_cnt  = count_sev("Critical")
        high_cnt      = count_sev("High")
        medium_cnt    = count_sev("Medium")
        low_cnt       = count_sev("Low")
        resolved_cnt  = 0  # placeholder – extend if you track resolved

        sc1, sc2, sc3, sc4, sc5, sc6 = st.columns(6)
        cards = [
            (sc1, "Total Risks",    total_risks,   "blue"),
            (sc2, "Critical",      critical_cnt,  "red"),
            (sc3, "High Severity", high_cnt,      "amber"),
            (sc4, "Moderate",      medium_cnt,    "amber"),
            (sc5, "Low Severity",  low_cnt,       "green"),
            (sc6, "Dependency Gaps", len(risk_data["dependencies"]), "purple"),
        ]
        for col, label, val, color in cards:
            with col:
                st.markdown(f"""
                <div class="stat-card">
                    <div class="stat-label">{label}</div>
                    <div class="stat-value {color}">{val}</div>
                </div>""", unsafe_allow_html=True)

        st.markdown("<div style='height:14px'></div>", unsafe_allow_html=True)

        # Filter pill buttons
        filter_opts = ["All", "Critical", "High", "Moderate", "Low"]
        fcols = st.columns(len(filter_opts) + 4)
        for i, f in enumerate(filter_opts):
            with fcols[i]:
                label_display = f
                if f != "All":
                    color_map_btn = {"Critical":"🔴","High":"🟠","Moderate":"🟡","Low":"🟢"}
                    label_display = f"{color_map_btn.get(f,'')} {f}"
                if st.button(label_display,
                             key=f"filter_pill_{f}",
                             type="primary" if st.session_state.risk_filter == f else "secondary",
                             width="stretch"):
                    st.session_state.risk_filter = f
                    st.rerun()

        # Chart columns
        chart_col, cat_col = st.columns([3, 2])
        with chart_col, st.container(border=True, key="card_risk_matrix"):
            severity_map = {"Moderate": "Medium", "All": "All"}
            sel = st.session_state.risk_filter
            chart_sev = severity_map.get(sel, sel)
            fig_mat = render_risk_matrix_chart(risks, severity_filter=chart_sev)
            if fig_mat:
                st.plotly_chart(fig_mat, width="stretch", key="risks_matrix_chart")
            else:
                st.info("No risks match the selected filter.")

        with cat_col, st.container(border=True, key="card_risk_cat"):
            fig_cat = render_risk_category_chart(risks)
            if fig_cat:
                st.plotly_chart(fig_cat, width="stretch", key="risks_cat_chart")
            else:
                st.info("No category data.")

        # ── Search + filter risk table ──
        st.markdown("<div style='height:6px'></div>", unsafe_allow_html=True)
        srch_col, fltr_col, _ = st.columns([3, 2, 3])
        with srch_col:
            search_term = st.text_input("🔍 Search risks by description, or mitigation…", key="risk_search", placeholder="Search risks…", label_visibility="collapsed")
        with fltr_col:
            sev_opts = ["All"] + sorted(df_risks["Severity"].unique().tolist()) if not df_risks.empty and "Severity" in df_risks.columns else ["All"]
            cat_filter = st.selectbox("Category: All", ["All"] + (sorted(df_risks["Risk Category"].unique().tolist()) if not df_risks.empty and "Risk Category" in df_risks.columns else []), key="risk_cat_filter", label_visibility="collapsed")

        if not df_risks.empty:
            display_df = df_risks.copy()
            if search_term:
                mask = display_df.apply(lambda r: search_term.lower() in str(r).lower(), axis=1)
                display_df = display_df[mask]
            if cat_filter != "All" and "Risk Category" in display_df.columns:
                display_df = display_df[display_df["Risk Category"] == cat_filter]
            st.dataframe(display_df, width="stretch", hide_index=True)
            st.caption(f"{len(display_df)} risk(s) shown · {len(df_risks)} total detected")
        else:
            st.success("No risks detected.")

        # Dependency gaps
        st.markdown("### 🔗 Dependency Gaps & Prerequisite Bottlenecks")
        if risk_data["dependencies"]:
            st.dataframe(pd.DataFrame(risk_data["dependencies"]), width="stretch", hide_index=True)
        else:
            st.info("No dependency gaps detected.")

        st.markdown("### 💡 Grounded Delivery Recommendations")
        for rec in forecast["recommendations"]:
            st.markdown(f"- {rec}")


# ══════════════════════════════════════════════════════════════
#  BLOCKERS PAGE
# ══════════════════════════════════════════════════════════════
elif page == "Blockers":
    page_header("🚧", "Blockers & Action Items",
                "Unresolved blockers, pending decisions, and assigned action items.")

    if not st.session_state.documents:
        empty_state("🚧", "No blockers or action items yet",
                    "Add meeting notes, status reports or trackers to find blockers, decisions and follow-ups.", "blk")
    else:
        blocker_data = st.session_state.blocker_agent.analyze(st.session_state.documents)
        b1, b2, b3, b4 = st.columns(4)
        b1.metric("Unresolved Blockers", blocker_data["summary"]["total_blockers"])
        b2.metric("Pending Decisions",   blocker_data["summary"]["total_decisions"])
        b3.metric("Total Action Items",  blocker_data["summary"]["total_action_items"])
        b4.metric("Open Action Items",   blocker_data["summary"]["open_action_items"])

        with st.container(border=True, key="card_blk_blockers"):
            section_title("🚨 Unresolved Blockers", "Issues currently stopping work, with severity and source.")
            show_table(blocker_data["blockers"], "No active blockers.", key="blk_blockers",
                       download_name="blockers", success_when_empty=True)
        with st.container(border=True, key="card_blk_decisions"):
            section_title("⚖️ Pending Decisions", "Open questions and approvals the team is waiting on.")
            show_table(blocker_data["decisions"], "No pending decisions.", key="blk_decisions", download_name="decisions")
        with st.container(border=True, key="card_blk_actions"):
            t_col, f_col = st.columns([3, 1])
            with t_col:
                section_title("✅ Action Item Tracker", "Follow-ups with assignee, due date and status.")
            df_act = pd.DataFrame(blocker_data["action_items"]) if blocker_data["action_items"] else pd.DataFrame()
            with f_col:
                statuses = sorted(df_act["Status"].dropna().astype(str).unique()) if "Status" in df_act.columns else []
                sf = st.selectbox("Filter by status", ["All"] + statuses, key="action_status_filter")
            if not df_act.empty:
                if sf != "All":
                    df_act = df_act[df_act["Status"].astype(str) == sf]
            show_table(df_act.to_dict("records"), "No action items detected.", key="blk_actions", download_name="action_items")


# ══════════════════════════════════════════════════════════════
#  HEALTH PAGE
# ══════════════════════════════════════════════════════════════
elif page == "Health":
    page_header("🩺", "Project Health", "Aggregate health score (0–100) with a dimension-by-dimension breakdown.")

    if not st.session_state.documents:
        empty_state("🩺", "No health score yet",
                    "Add project documents to calculate scope, timeline and blocker health.", "health")
    else:
        scope_data   = st.session_state.scope_agent.analyze(st.session_state.documents)
        risk_data    = st.session_state.risk_agent.analyze(st.session_state.documents, scope_data["milestones"])
        blocker_data = st.session_state.blocker_agent.analyze(st.session_state.documents)
        health_data  = st.session_state.health_scorer.calculate_health_score(st.session_state.documents, scope_data, risk_data, blocker_data)

        overall    = health_data["overall_score"]
        badge_cls  = health_data["badge_class"]
        dim        = health_data["dimension_scores"]

        c1, c2 = st.columns([2, 3])
        with c1:
            with st.container(border=True, key="card_health_overall"):
                section_title("Overall project health")
                st.plotly_chart(health_gauge(overall, 230), width="stretch",
                                key="health_gauge", config={"displayModeBar": False})
                st.markdown(f"<div style='text-align:center;'><span class='badge {badge_cls}' style='font-size:14px; padding:6px 18px;'>"
                            f"{health_data['health_status']}</span></div>", unsafe_allow_html=True)
                st.caption(health_data["summary_description"])
        with c2:
            with st.container(border=True, key="card_health_dims"):
                section_title("Dimension scores", "0–100, higher is healthier.")
                score_bar("Scope clarity", dim["scope_clarity"], "30% of overall — how clearly goals, milestones, deliverables and owners are defined.")
                score_bar("Timeline", dim["timeline_risk"], "40% of overall — schedule risk from milestone dates, delays and risk signals.")
                score_bar("Blockers", dim["blocker_impact"], "30% of overall — impact of unresolved blockers, decisions and open actions.")

        d1, d2 = st.columns(2)
        with d1:
            with st.container(border=True, key="card_health_pos"):
                section_title("🌟 Positive drivers")
                all_dr = health_data["positive_drivers"]["scope"] + health_data["positive_drivers"]["timeline"] + health_data["positive_drivers"]["blockers"]
                for d in all_dr:
                    st.markdown(f"- ✅ {d}")
                if not all_dr:
                    st.info("No positive drivers.")
        with d2:
            with st.container(border=True, key="card_health_neg"):
                section_title("⚠️ Penalties & risk drivers")
                all_pe = health_data["penalties"]["scope"] + health_data["penalties"]["timeline"] + health_data["penalties"]["blockers"]
                for p in all_pe:
                    st.markdown(f"- 🚩 {p}")
                if not all_pe:
                    st.success("No penalties.")
        with st.container(border=True, key="card_health_recs"):
            section_title("💡 Actionable guidance")
            for r in health_data["recommendations"]:
                st.markdown(f"- {r}")


# ══════════════════════════════════════════════════════════════
#  AI ASSISTANT PAGE
# ══════════════════════════════════════════════════════════════
elif page == "AIAssistant":
    page_header("💬", "AI Assistant", "Conversational assistant grounded in your project documents, with sources for every answer.")

    if not st.session_state.documents:
        empty_state("💬", "Nothing to chat about yet",
                    "Add project documents and the assistant will answer questions using only their content.", "ai")
    else:
        top_l, top_r = st.columns([5, 1])
        with top_l:
            st.caption("Answers use only your uploaded documents and cite their sources. "
                       "Questions the documents can't answer are declined rather than guessed.")
        with top_r:
            if st.button("Clear chat", icon="🧹", width="stretch", disabled=not st.session_state.chat_history):
                st.session_state.chat_history = []
                st.rerun()

        chip_c1, chip_c2, chip_c3, chip_c4, chip_c5 = st.columns(5)
        selected_prompt = None
        if chip_c1.button("🟢 Are we on track?",    width="stretch"): selected_prompt = "Are we on track?"
        if chip_c2.button("🚩 Biggest risks?",       width="stretch"): selected_prompt = "What are our biggest risks?"
        if chip_c3.button("🚨 Active blockers?",     width="stretch"): selected_prompt = "What active blockers are open?"
        if chip_c4.button("👥 Deliverable RACI?",    width="stretch"): selected_prompt = "Who is responsible for deliverables?"
        if chip_c5.button("✅ Action items?",         width="stretch"): selected_prompt = "What action items are assigned?"

        if not st.session_state.chat_history:
            with st.chat_message("assistant"):
                st.markdown(
                    f"Hi! I've read **{len(st.session_state.indexed_sources)} document(s)** "
                    f"({len(st.session_state.chunks)} indexed chunks). Ask me about schedule, risks, blockers, "
                    "decisions, owners or action items — or tap one of the suggestions above."
                )

        for msg in st.session_state.chat_history:
            with st.chat_message(msg["role"]):
                st.markdown(msg["content"])
                if msg.get("sources"):
                    with st.expander(f"📌 Grounded Evidence ({len(msg['sources'])} sources)"):
                        for s in msg["sources"]:
                            st.markdown(f"• **{s.get('Source')}** ({s.get('Location')}) — {s.get('Reason', s.get('Score',''))}")

        chat_input = st.chat_input("Ask a question about your project…")
        user_query = selected_prompt or chat_input
        if user_query:
            st.session_state.chat_history.append({"role": "user", "content": user_query})
            with st.chat_message("user"):
                st.markdown(user_query)

            with st.spinner("Searching your documents…"):
                scope_data   = st.session_state.scope_agent.analyze(st.session_state.documents)
                risk_data    = st.session_state.risk_agent.analyze(st.session_state.documents, scope_data["milestones"])
                blocker_data = st.session_state.blocker_agent.analyze(st.session_state.documents)
                health_data  = st.session_state.health_scorer.calculate_health_score(st.session_state.documents, scope_data, risk_data, blocker_data)

                response = st.session_state.conversational_assistant.answer_query(
                    query=user_query, records=st.session_state.documents,
                    chunks=st.session_state.chunks, vectorizer=st.session_state.vectorizer,
                    matrix=st.session_state.matrix, top_k=TOP_K,
                    scope_data=scope_data, risk_data=risk_data,
                    blocker_data=blocker_data, health_data=health_data,
                )
            sources = response.get("sources", [])
            st.session_state.chat_history.append({"role": "assistant", "content": response["answer"], "sources": sources})
            st.rerun()


# ══════════════════════════════════════════════════════════════
#  DOC GEN PAGE
# ══════════════════════════════════════════════════════════════
elif page == "DocGen":
    page_header("📄", "Reports & Doc Generation",
                "User stories, risk register, and action item list generated from your project content.")

    if not st.session_state.documents:
        empty_state("📄", "No reports yet",
                    "Add project documents to generate user stories, a risk register and an action item list.", "dg")
    else:
        doc_data = st.session_state.doc_gen_agent.analyze(st.session_state.documents)
        m1, m2, m3, m4 = st.columns([1, 1, 1, 2])
        m1.metric("User Stories",          len(doc_data["user_stories"]))
        m2.metric("Risk Register Entries", len(doc_data["risk_register"]))
        m3.metric("Action Items",          len(doc_data["action_items"]))
        with m4:
            st.markdown("<div style='height:18px'></div>", unsafe_allow_html=True)
            st.download_button("Download full governance report (Markdown)",
                               doc_data["full_report"], "project_intelligence_report.md",
                               "text/markdown", type="primary", icon="📥", width="stretch")
        with st.container(border=True, key="card_dg_stories"):
            section_title("🎯 User Stories", "Backlog-ready stories with acceptance criteria.")
            show_table(doc_data["user_stories"], "No user stories.", key="dg_stories", download_name="user_stories")
        with st.container(border=True, key="card_dg_register"):
            section_title("⚠️ Risk Register", "Each risk with severity, owner and mitigation.")
            show_table(doc_data["risk_register"], "No risk register entries.", key="dg_register", download_name="risk_register")
        with st.container(border=True, key="card_dg_actions"):
            section_title("🚧 Action Items", "Owned follow-ups to drive the project forward.")
            show_table(doc_data["action_items"], "No action items.", key="dg_actions", download_name="action_items_report")
        with st.expander("👁️ Preview the governance report"):
            st.markdown(doc_data["full_report"])


# ══════════════════════════════════════════════════════════════
#  ARCHITECTURE & MILESTONES PAGE
# ══════════════════════════════════════════════════════════════
elif page == "Architecture":
    st.markdown("""
    <div class="page-header-bar">
        <div>
            <div class="page-title">🗺️ Architecture & Milestones</div>
            <div class="page-sub">RAG index explorer, retrieval engine, and milestone delivery tracking.</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    st.markdown("#### Milestone 4 — Insights, Incremental Knowledge Base, Testing")
    st.markdown("""
    1. **Project Insights & Risk Summary Dashboard** — health score, extracted risks, action items, and scope summary.
    2. **Incremental document upload** — new meeting notes / progress updates merge into the knowledge base and rebuild the RAG index automatically.
    3. **End-to-end testing** — Settings → Run End-to-End Validation, plus `test_agents.py`, `test_conversational_rag.py`, and `test_milestone4.py`.
    4. **Documentation** — see `README.md` and `PROJECT_REPORT.md` for architecture, demo script, and evaluation mapping.
    """)
    st.divider()
    st.markdown("#### 🧩 RAG Index")
    a, b, c = st.columns(3)
    a.metric("Chunk Size",    CHUNK_SIZE)
    b.metric("Chunk Overlap", CHUNK_OVERLAP)
    c.metric("Indexed Chunks",len(st.session_state.chunks))

    if st.button("🔄 Rebuild RAG Index"):
        st.session_state.chunks = build_chunks(st.session_state.documents)
        rebuild_index(st.session_state.chunks)
        st.success("Index rebuilt.")

    if st.session_state.chunks:
        preview_df = pd.DataFrame([{
            "Chunk ID": ch.chunk_id, "Source": ch.source,
            "Location": ch.location, "Chars": len(ch.text),
            "Preview": ch.text[:150],
        } for ch in st.session_state.chunks[:50]])
        st.dataframe(preview_df, width="stretch", hide_index=True)

    st.divider()
    st.markdown("#### 🔎 Retrieval Explorer")
    q = st.text_input("Enter a project question", placeholder="What are the current project risks?", key="arch_query")
    if st.button("Retrieve Chunks", type="primary"):
        if not q.strip(): st.warning("Enter a question.")
        else:
            results = retrieve(q, TOP_K)
            for r in results:
                ch = r["chunk"]
                st.markdown(f"**{ch.chunk_id} — {ch.source} — {ch.location}** | Score: `{r['score']:.3f}`")
                st.write(ch.text)
                st.divider()


# ══════════════════════════════════════════════════════════════
#  SETTINGS PAGE
# ══════════════════════════════════════════════════════════════
elif page == "Settings":
    st.markdown("""
    <div class="page-header-bar">
        <div>
            <div class="page-title">⚙️ Settings</div>
            <div class="page-sub">RAG configuration, ingestion parameters, and validation suite.</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    st.markdown("#### RAG Parameters")
    chunk_size    = st.slider("Chunk Size (characters)",   200, 1500, CHUNK_SIZE,    50)
    chunk_overlap = st.slider("Chunk Overlap",              0,   500, CHUNK_OVERLAP,  10)
    top_k         = st.slider("Retrieval Top-K",            1,    10, TOP_K)

    st.divider()
    st.markdown("#### 🧪 Validation Suite")
    if st.button("▶️ Run End-to-End Validation", type="primary"):
        from test_conversational_rag import MockUploadedFile
        sample_dir = Path("sample_data")
        if not sample_dir.exists():
            st.error("sample_data directory not found.")
        else:
            files = list(sample_dir.glob("*.*"))
            test_records, file_results = [], []
            for fpath in files:
                mock = MockUploadedFile(str(fpath))
                try:
                    recs = extract_file(mock)
                    test_records.extend(recs)
                    file_results.append({"Document": mock.name, "Format": fpath.suffix.upper(), "Status": "PASSED", "Units": len(recs)})
                except Exception as exc:
                    file_results.append({"Document": mock.name, "Format": fpath.suffix.upper(), "Status": f"FAILED: {exc}", "Units": 0})

            st.markdown("**File Format Ingestion:**")
            st.dataframe(pd.DataFrame(file_results), width="stretch", hide_index=True)

            scope_r   = st.session_state.scope_agent.analyze(test_records)
            risk_r    = st.session_state.risk_agent.analyze(test_records, scope_r["milestones"])
            blocker_r = st.session_state.blocker_agent.analyze(test_records)
            doc_r     = st.session_state.doc_gen_agent.analyze(test_records)
            health_r  = st.session_state.health_scorer.calculate_health_score(test_records, scope_r, risk_r, blocker_r)

            val = [
                {"Agent": "Scope Agent",       "Metric": "Goals",          "Count": str(len(scope_r["goals"])),       "Status": "PASSED" if scope_r["goals"] else "FAILED"},
                {"Agent": "Risk Agent",         "Metric": "Risks",          "Count": str(len(risk_r["risks"])),        "Status": "PASSED" if risk_r["risks"] else "FAILED"},
                {"Agent": "Blocker Agent",      "Metric": "Blockers",       "Count": str(len(blocker_r["blockers"])),  "Status": "PASSED" if blocker_r["blockers"] else "FAILED"},
                {"Agent": "Doc Gen Agent",      "Metric": "User Stories",   "Count": str(len(doc_r["user_stories"])),  "Status": "PASSED" if doc_r["user_stories"] else "FAILED"},
                {"Agent": "Health Scorer",      "Metric": "Overall Score",  "Count": f"{health_r['overall_score']}/100", "Status": f"PASSED ({health_r['health_status']})"},
            ]
            st.markdown("**Multi-Agent Validation Matrix:**")
            st.dataframe(pd.DataFrame(val), width="stretch", hide_index=True)

            update_path = sample_dir / "progress_update.txt"
            if update_path.exists():
                incr = extract_file(MockUploadedFile(str(update_path)))
                merged = merge_documents_incremental(test_records, incr)
                st.markdown("**Incremental knowledge-base update:**")
                st.write(
                    f"Baseline units: {len(test_records)} → after progress update: {len(merged)} "
                    f"(+{len(merged) - len(test_records)}). RAG chunks would rebuild automatically on ingest."
                )
                assert_ok = len(merged) > len(test_records)
                if assert_ok:
                    st.success("Incremental merge added new units.")
                else:
                    st.error("Incremental merge did not grow the knowledge base.")

            st.success("🎉 All validation tests completed!")


